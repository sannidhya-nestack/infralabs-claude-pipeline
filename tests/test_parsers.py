"""Parser contract tests on synthetic bills. Every number and name here is made up: this repo is
public, so a fixture never holds a line from a real bill."""
import pathlib, sys
from decimal import Decimal as D

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "scripts" / "parsers"))

import common  # noqa: E402
import detect  # noqa: E402
import optus_bill_to_xlsx as optus  # noqa: E402

PAGE1 = """\
                                                            Page 1 of 3
Optus Billing Services Pty. Ltd.                   Invoice No:      000000000001
                                                   ACCOUNT NUMBER      11111111111
    EXAMPLE TRADING PTY LTD
Issue Date: 01 Mar 2026
Invoice Period: 01 Feb 2026 to 28 Feb 2026
  New Charges                     $56.50
  Recurring Charges               $40.00
  Usage Charges                   $12.00
  Other Charges and Credits        $0.00
  GST                              $4.50
"""
PAGE2 = """\
Service Details
    0400000001   Total Charges (Ex GST)   $30.00
    Plan A
    Jane Citizen
  Recurring Charges     Description          Total   $25.00
  Monthly plan fee      01 Feb - 28 Feb              $25.00
  Usage Charges         Description          Total    $5.00
  Calls                 10 calls                      $5.00
Service Details
    0400000002   Total Charges (Ex GST)   $22.00
    Plan B
    John Citizen
  Recurring Charges     Description          Total   $15.00
  Monthly plan fee      01 Feb - 28 Feb              $20.00
  Loyalty discount      01 Feb - 28 Feb         credit 5.00
"""
PAGE3 = """\
Service Details (continued)
    0400000002   Total Charges (Ex GST)   $22.00
    Plan B
    John Citizen
  Usage Charges         Description          Total    $7.00
  *Roaming data         1 GB                          $7.00
Usage Details
  0400000002  01 Feb  Data  1 GB                       $7.00
"""
PAGES = [PAGE1, PAGE2, PAGE3]


def test_synthetic_optus_bill_passes():
    parsed = optus.parse(PAGES)
    r = common.report(parsed, D("1.00"))
    assert r["status"] == "PASS", r
    assert r["rows"] == 2
    assert r["total_inc_gst"]["difference"] == "0.00"
    # 20 - 5 credit at 10% GST, 7 roaming GST-free
    assert parsed.rows[1][4:7] == [16.5, 7.0, 23.5]


def test_a_dropped_service_fails():
    # Without its header, service 1 is never read: page 1's recurring total no longer matches.
    page2 = PAGE2.replace("    0400000001   Total Charges (Ex GST)   $30.00\n", "")
    r = common.report(optus.parse([PAGE1, page2, PAGE3]), D("1.00"))
    assert r["status"] == "FAIL"
    assert r["error_count"] > 0


def test_detect():
    assert detect.detect(PAGE1) == ("digital", "optus")
    assert detect.detect("   \n  ") == ("scanned", "none")
    assert detect.detect("Some Other Carrier Pty Ltd  Tax Invoice  Account 123  Total $10.00\n" * 2) == ("digital", "none")


def test_workbook(tmp_path):
    from openpyxl import load_workbook

    out = tmp_path / "w.xlsx"
    common.write_xlsx(out, optus.parse(PAGES))
    wb = load_workbook(out)
    assert wb["Summary"]["C6"].value == 56.5
    assert wb["Summary"]["H13"].value == 0.0
    sheet = wb["11111111111"]
    assert sheet["A1"].value == "Identifiying #"
    assert sheet["O1"].value == "Calls to 13 Numbers - Discount (count)"
    assert sheet["H2"].value == "Feb"
    assert [sheet.cell(r, 1).value for r in (3, 4)] == ["0400000001", "0400000002"]
